from datetime import datetime
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt, Inches


def add_title(doc, text, size=22):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)


def add_subtitle(doc, text, size=13):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_body(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style='List Bullet')
        for run in p.runs:
            run.font.size = Pt(11)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style='List Number')
        for run in p.runs:
            run.font.size = Pt(11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    return table


def build_document(output_path):
    doc = Document()

    # Margins
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # Cover page
    add_title(doc, 'SISTEMA DE INVENTARIO Y GESTION TECNICA', 24)
    add_title(doc, 'Taller de Electronica - Armada de Republica Dominicana', 16)
    doc.add_paragraph('')
    add_subtitle(doc, 'Documentacion Tecnica y Funcional para Presentacion de Monografico')
    doc.add_paragraph('')
    add_subtitle(doc, f'Fecha: {datetime.now().strftime("%d/%m/%Y")}', 12)
    add_subtitle(doc, 'Version del sistema: 1.0.0', 12)
    add_subtitle(doc, 'Stack: React + Django REST + PostgreSQL', 12)

    doc.add_page_break()

    # Executive summary
    add_heading(doc, '1. Resumen Ejecutivo', level=1)
    add_body(
        doc,
        'Este documento presenta la arquitectura, diseno, implementacion, seguridad, despliegue y operacion '
        'del Sistema de Inventario del Taller de Electronica de la Armada de Republica Dominicana. '
        'El objetivo principal es garantizar trazabilidad completa de inventario, ordenes tecnicas, custodia de herramientas '
        'y auditoria institucional con controles de acceso por rol.'
    )

    add_heading(doc, '2. Objetivos del Proyecto', level=1)
    add_bullets(doc, [
        'Centralizar la gestion de inventario y operaciones tecnicas en una unica plataforma.',
        'Mejorar el control de entradas, salidas y stock critico en componentes esenciales para la flota.',
        'Digitalizar las ordenes de servicio de mantenimiento, reparacion e instalacion.',
        'Asegurar transparencia mediante auditoria inmutable de operaciones sensibles.',
        'Aplicar seguridad por roles y permisos dinamicos en frontend y backend.'
    ])

    add_heading(doc, '3. Alcance Funcional', level=1)
    add_bullets(doc, [
        'Gestion de usuarios, roles y permisos (RBAC).',
        'Inventario: productos, ubicaciones, movimientos, recepciones y reportes.',
        'Ordenes de servicio: ciclo de vida, tecnicos asignados y cierre tecnico.',
        'Despachos y solicitantes: salidas de materiales con comprobantes.',
        'Custodia de herramientas: unidades por serial, prestamos, devoluciones y estado.',
        'Auditoria institucional con trazabilidad por usuario, IP, accion y fecha.'
    ])

    add_heading(doc, '4. Arquitectura de la Solucion', level=1)
    add_body(doc, 'La solucion implementa una arquitectura cliente-servidor desacoplada:')
    add_bullets(doc, [
        'Frontend SPA: React 18 + Vite + TailwindCSS.',
        'Backend API REST: Django + Django REST Framework.',
        'Autenticacion: JWT (access/refresh) con timeout de sesion.',
        'Base de datos: PostgreSQL en produccion (SQLite en desarrollo local).',
        'Contenedores: Docker Compose para entorno de desarrollo y produccion.'
    ])

    add_heading(doc, '5. Stack Tecnologico y Dependencias', level=1)
    add_heading(doc, '5.1 Backend', level=2)
    add_table(doc,
              ['Componente', 'Detalle'],
              [
                  ['Framework', 'Django >=5.0,<6.0'],
                  ['API', 'Django REST Framework'],
                  ['Auth', 'djangorestframework-simplejwt'],
                  ['Filtros', 'django-filter'],
                  ['OpenAPI', 'drf-spectacular'],
                  ['DB Adapter', 'psycopg2-binary'],
                  ['Reportes', 'reportlab, openpyxl'],
                  ['WSGI', 'gunicorn, whitenoise'],
              ])

    add_heading(doc, '5.2 Frontend', level=2)
    add_table(doc,
              ['Componente', 'Detalle'],
              [
                  ['Framework', 'React 18'],
                  ['Build Tool', 'Vite 5'],
                  ['UI', 'TailwindCSS + Headless UI + Heroicons'],
                  ['HTTP Client', 'Axios'],
                  ['Routing', 'react-router-dom'],
                  ['Notificaciones', 'react-hot-toast'],
                  ['Graficos', 'recharts'],
              ])

    add_heading(doc, '6. Estructura del Proyecto', level=1)
    add_bullets(doc, [
        'backend/accounts: autenticacion, usuarios, settings de sesion, permisos por rol.',
        'backend/inventory: inventario, catalogos, movimientos, recepcion, unidades y prestamos.',
        'backend/workorders: ordenes de servicio, despachos y solicitantes.',
        'backend/audit: auditoria y trazabilidad de operaciones.',
        'frontend/src/pages: vistas operativas por modulo.',
        'frontend/src/services: integracion HTTP con API.',
        'frontend/src/context: estado global de autenticacion.',
        'frontend/src/config: rutas, menu y navegacion condicionada por permisos.'
    ])

    add_heading(doc, '7. Modelo de Seguridad y Control de Acceso', level=1)
    add_body(doc, 'El sistema aplica seguridad en capas: autenticacion JWT, autorizacion por rol y matriz de permisos dinamicos.')
    add_heading(doc, '7.1 Roles institucionales', level=2)
    add_bullets(doc, [
        'Administrador',
        'Encargado de Inventario (Almacenista)',
        'Tecnico'
    ])

    add_heading(doc, '7.2 Matriz de permisos dinamicos', level=2)
    add_bullets(doc, [
        'users.manage', 'roles.manage',
        'inventory.view', 'inventory.manage',
        'products.view', 'products.manage',
        'service_orders.view', 'service_orders.manage',
        'reception.view', 'reception.manage',
        'despachos.view', 'despachos.manage',
        'solicitantes.view', 'solicitantes.manage',
        'locations.view', 'locations.manage',
        'catalogs.view', 'catalogs.manage',
        'audit.view', 'reports.export'
    ])

    add_heading(doc, '7.3 Blindaje backend', level=2)
    add_body(
        doc,
        'La API valida permisos en cada endpoint. Aunque un usuario intente llamadas manuales, '
        'las operaciones quedan restringidas por la matriz de permisos del rol autenticado.'
    )

    add_heading(doc, '8. Modulos Funcionales', level=1)
    add_heading(doc, '8.1 Inventario', level=2)
    add_bullets(doc, [
        'CRUD de productos y catalogos.',
        'Movimientos de stock (entrada/salida).',
        'Control de stock minimo y estado critico.',
        'Ubicacion fisica por nomenclatura militar.',
        'Reportes PDF/Excel de inventario y stock critico.'
    ])

    add_heading(doc, '8.2 Recepcion de Mercancia', level=2)
    add_bullets(doc, [
        'Recepcion por lotes con reception_id unico.',
        'Registro de documento, entregado por, cedula y rango/cargo.',
        'Historial filtrable y exportable.',
        'Comprobante final con seriales de equipos/partes.'
    ])

    add_heading(doc, '8.3 Despachos', level=2)
    add_bullets(doc, [
        'Despacho atomico con descuento de stock y trazabilidad.',
        'Relacion con solicitante y unidad.',
        'Anulacion controlada con motivo y reversion.',
        'Comprobante de despacho con serial de unidad/equipo.'
    ])

    add_heading(doc, '8.4 Ordenes de Servicio', level=2)
    add_bullets(doc, [
        'Flujo de estados: recibido, diagnostico, proceso, completado, entregado, cancelado.',
        'Asignacion de tecnico y bitacora de eventos.',
        'Notas tecnicas y cierre con diagnostico/trabajo realizado.',
        'Comprobante de cierre tecnico en PDF/Excel.'
    ])

    add_heading(doc, '8.5 Herramientas y Custodia', level=2)
    add_bullets(doc, [
        'Unidades por serial y estado operativo.',
        'Prestamos y devoluciones con control de vencimiento.',
        'Estados: disponible, asignado, mantenimiento, dado de baja.'
    ])

    add_heading(doc, '8.6 Auditoria', level=2)
    add_bullets(doc, [
        'Registro de acciones create/update/delete.',
        'Atributos auditados: usuario, accion, timestamp, IP, modelo y objeto.',
        'Consulta de auditoria protegida por permiso audit.view.'
    ])

    add_heading(doc, '9. API REST (Base: /api/v1/)', level=1)
    add_heading(doc, '9.1 Autenticacion y cuentas', level=2)
    add_bullets(doc, [
        'POST /auth/login/',
        'POST /auth/refresh/',
        'POST /auth/logout/',
        'GET /auth/me/',
        'POST /auth/change-password/',
        'POST /auth/admin-reset-password/',
        'GET/POST /users/',
        'GET/PATCH/PUT /role-permissions/'
    ])

    add_heading(doc, '9.2 Inventario y productos', level=2)
    add_bullets(doc, [
        'GET/POST /inventory/items/',
        'GET /inventory/items/report/?type=pdf|excel',
        'GET /inventory/items/critical/',
        'POST /inventory/items/{id}/add_unit/',
        'GET/POST /inventory/product-entries/',
        'GET /inventory/product-entries/receipt/',
        'GET /inventory/product-entries/history-report/',
        'GET/POST /products/items/'
    ])

    add_heading(doc, '9.3 Work Orders', level=2)
    add_bullets(doc, [
        'GET/POST /work-orders/service-orders/',
        'POST /work-orders/service-orders/{id}/assign/',
        'POST /work-orders/service-orders/{id}/transition/',
        'POST /work-orders/service-orders/{id}/add_note/',
        'POST /work-orders/service-orders/{id}/complete_service/',
        'GET /work-orders/service-orders/{id}/completion_receipt/',
        'GET/POST /work-orders/despachos/',
        'GET /work-orders/despachos/{id}/receipt/'
    ])

    add_heading(doc, '10. Base de Datos (Resumen de Entidades)', level=1)
    add_table(doc,
              ['Modulo', 'Entidades principales'],
              [
                  ['Accounts', 'User, RolePermission, SystemSetting'],
                  ['Inventory', 'Item, Category, Location, StockMovement, Transfer, ItemUnit, ItemLoan, EntradaProducto'],
                  ['Workorders', 'ServiceOrder, ServiceOrderLog, Despacho, LineaDespacho, Solicitante'],
                  ['Audit', 'AuditLog']
              ])

    add_heading(doc, '11. Reportes y Evidencias', level=1)
    add_bullets(doc, [
        'Reportes PDF/Excel de inventario, stock critico y herramientas.',
        'Historial de recepcion exportable con filtros.',
        'Comprobantes de recepcion, despacho y cierre de orden.',
        'Reportes recepcion vs despacho para control logistico.'
    ])

    add_heading(doc, '12. Despliegue y Operacion', level=1)
    add_heading(doc, '12.1 Entorno local con Docker', level=2)
    add_numbered(doc, [
        'docker compose up -d --build --wait',
        'Verificar estado con docker compose ps',
        'Backend: /api/v1/docs/ para validar endpoints'
    ])

    add_heading(doc, '12.2 Entorno de produccion', level=2)
    add_numbered(doc, [
        'Crear archivo .env.prod con secretos y hosts reales.',
        'Ejecutar docker compose --env-file .env.prod -f docker-compose.prod.yml up -d --build',
        'Aplicar migraciones: docker compose --env-file .env.prod -f docker-compose.prod.yml run --rm backend python manage.py migrate',
        'Validar salud con /health/ y logs de servicios.'
    ])

    add_heading(doc, '13. Calidad, Pruebas y Validacion', level=1)
    add_bullets(doc, [
        'Validacion backend con python manage.py check.',
        'Build frontend con npm run build.',
        'Pruebas funcionales por rol (admin, almacenista, tecnico).',
        'Pruebas de seguridad de rutas y endpoints con token valido/no valido.'
    ])

    add_heading(doc, '14. Riesgos Tecnicos y Mitigaciones', level=1)
    add_table(doc,
              ['Riesgo', 'Impacto', 'Mitigacion'],
              [
                  ['Permisos mal configurados', 'Acceso no autorizado', 'Matriz dinamica centralizada + validacion backend'],
                  ['Fallo de despliegue en VM', 'Interrupcion de servicio', 'Healthchecks, logs y scripts de arranque'],
                  ['Datos incompletos en recepcion', 'Trazabilidad parcial', 'Validaciones de formulario y reglas de negocio'],
                  ['Stock inconsistente', 'Errores operativos', 'Operaciones atomicas con transacciones DB']
              ])

    add_heading(doc, '15. Conclusion', level=1)
    add_body(
        doc,
        'El sistema implementado cumple con los requerimientos academicos y operativos de una plataforma '
        'institucional moderna: seguridad por capas, trazabilidad total, reportabilidad formal y operacion '
        'en entornos locales o de produccion. La arquitectura propuesta es escalable y mantenible para futuras '
        'integraciones logisticas de la Armada.'
    )

    add_heading(doc, '16. Anexos', level=1)
    add_bullets(doc, [
        'Manual de Usuario (MANUAL_USUARIO.md).',
        'Especificacion funcional (SPEC.md).',
        'Guia de diseno (DESIGN.md).',
        'Documentacion de despliegue Oracle/Produccion (Oracle.txt).'
    ])

    doc.save(output_path)


if __name__ == '__main__':
    build_document('c:/Projecto/SistemadeInventario-ARD/Documentacion_Monografico_SistemaInventario_ARD.docx')
    print('Documento generado correctamente.')
