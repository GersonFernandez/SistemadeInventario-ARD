import io
from datetime import datetime
from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter, column_index_from_string


def _resolve_logo_path():
    """Resolve institutional logo path from known project locations."""
    backend_root = Path(__file__).resolve().parents[1]
    repo_root = backend_root.parent
    candidates = [
        backend_root / 'logo.jpeg',
        backend_root / 'logo.jpg',
        backend_root / 'logo.png',
        repo_root / 'logo.jpeg',
        repo_root / 'logo.jpg',
        repo_root / 'logo.png',
        repo_root / 'frontend' / 'src' / 'assets' / 'logo.jpeg',
        repo_root / 'frontend' / 'src' / 'assets' / 'logo.png',
    ]
    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return str(candidate)
    return None


def _fit_box(width, height, max_width, max_height):
    """Scale dimensions to fit inside a bounding box while preserving aspect ratio."""
    if not width or not height:
        return max_width, max_height
    scale = min(max_width / float(width), max_height / float(height))
    return width * scale, height * scale


def _cell_text(value):
    if value is None:
        return ''
    return str(value)


def _compute_pdf_col_widths(headers, rows):
    """Calculate widths that adapt to the real content without truncating text."""
    if not headers:
        return []

    max_lengths = [len(_cell_text(header)) for header in headers]
    for row in rows:
        for index, value in enumerate(row):
            if index >= len(max_lengths):
                break
            max_lengths[index] = max(max_lengths[index], len(_cell_text(value)))

    widths = [max(80, min(260, length * 7.5 + 26)) for length in max_lengths]
    page_width = 740
    total_width = sum(widths)
    if total_width > page_width:
        factor = page_width / float(total_width)
        widths = [max(60, width * factor) for width in widths]
        while sum(widths) > page_width and any(width > 60 for width in widths):
            for index in range(len(widths) - 1, -1, -1):
                if widths[index] > 60:
                    widths[index] = max(60, widths[index] - 10)
                    break
            else:
                break
    return widths


def _build_pdf_response(
    title,
    headers,
    rows,
    include_signatures=False,
    signature_names=None,
    metadata_lines=None,
    receipt_mode=False,
):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=28,
        rightMargin=28,
        topMargin=20,
        bottomMargin=18,
    )
    styles = getSampleStyleSheet()
    elements = []

    heading_style = ParagraphStyle(
        'HeadingReport',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#16365d'),
        alignment=TA_CENTER,
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        'SubtitleReport',
        parent=styles['Normal'],
        fontSize=9,
        textColor=colors.HexColor('#4b5563'),
        alignment=TA_CENTER,
        spaceAfter=10,
    )
    meta_label_style = ParagraphStyle(
        'MetaLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        textColor=colors.HexColor('#1f2937'),
        alignment=TA_LEFT,
        leading=12,
    )
    meta_value_style = ParagraphStyle(
        'MetaValue',
        parent=styles['Normal'],
        fontSize=8.5,
        textColor=colors.HexColor('#374151'),
        alignment=TA_LEFT,
        leading=12,
    )
    section_style = ParagraphStyle(
        'SectionTitle',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        textColor=colors.HexColor('#1f2937'),
        alignment=TA_CENTER,
        spaceAfter=6,
    )

    logo_path = _resolve_logo_path()
    header_logo = None
    if logo_path:
        try:
            header_logo = Image(logo_path)
            fitted_width, fitted_height = _fit_box(
                getattr(header_logo, 'imageWidth', 0),
                getattr(header_logo, 'imageHeight', 0),
                min(doc.width, 300),
                110,
            )
            header_logo.drawWidth = fitted_width
            header_logo.drawHeight = fitted_height
            header_logo.hAlign = 'CENTER'
        except Exception:
            header_logo = None

    if header_logo:
        header_logo.hAlign = 'CENTER'
        elements.append(header_logo)
        elements.append(Spacer(1, 8))
        elements.append(Paragraph('''<hr width="90%" color="#16365d" size="1"/>''', subtitle_style))
        elements.append(Spacer(1, 8))

    elements.append(Paragraph(title, heading_style))
    elements.append(Paragraph('''<hr width="88%" color="#16365d" size="1"/>''', subtitle_style))
    if not receipt_mode:
        elements.append(Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))

    if metadata_lines:
        elements.append(Spacer(1, 12))
        elements.append(Paragraph('Datos del comprobante', section_style))
        metadata_rows = []
        for line in metadata_lines:
            if ': ' in line:
                label, value = line.split(': ', 1)
            else:
                label, value = line, ''
            metadata_rows.append([
                Paragraph(label, meta_label_style),
                Paragraph(value if value else '—', meta_value_style),
            ])
        metadata_table = Table(metadata_rows, colWidths=[220, 460], hAlign='CENTER')
        metadata_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dbe2ea')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ]))
        elements.append(metadata_table)

    elements.append(Spacer(1, 16))
    elements.append(Paragraph('Detalle de equipos / piezas', section_style))

    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['BodyText'],
        fontSize=8.2,
        leading=10,
        alignment=TA_CENTER,
        wordWrap='CJK',
    )
    data = [[Paragraph(_cell_text(header), table_cell_style) for header in headers]]
    data.extend([
        [Paragraph(_cell_text(value), table_cell_style) for value in row]
        for row in rows
    ])
    col_widths = _compute_pdf_col_widths(headers, rows)
    table = Table(data, colWidths=col_widths or None, repeatRows=1, hAlign='CENTER', splitByRow=1, rowHeights=None)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#16365d')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#ffffff')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.HexColor('#ffffff')]),
        ('GRID', (0, 0), (-1, -1), 0.8, colors.HexColor('#cbd5e1')),
        ('FONTSIZE', (0, 1), (-1, -1), 8.5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 7),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ('LINEABOVE', (0, 1), (-1, 1), 1.2, colors.HexColor('#dbe2ea')),
    ]))
    elements.append(table)

    if include_signatures:
        elements.append(Spacer(1, 22))
        delivered_name = (signature_names or {}).get('delivered_by', '')
        received_name = (signature_names or {}).get('received_by', '')
        signature_headers = ['Entregado por', 'Recibido por']
        signature_rows = [[delivered_name, received_name], ['______________________________', '______________________________']]
        sig_table = Table([signature_headers] + signature_rows, colWidths=[280, 280], hAlign='CENTER')
        sig_table.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eef2ff')),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        elements.append(sig_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def _build_excel_response(title, headers, rows, receipt_mode=False):
    buffer = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    # Excel no permite estos caracteres en el nombre del sheet
    safe_title = title.replace('/', '-').replace('\\', '-').replace(':', '-').replace('*', '-').replace('?', '-').replace('[', '-').replace(']', '-')[:31]
    ws.title = safe_title

    logo_path = _resolve_logo_path()
    first_title_row = 1
    title_start_column = 'A'
    if logo_path:
        try:
            img = XLImage(logo_path)
            fitted_width, fitted_height = _fit_box(img.width, img.height, 240, 80)
            img.width = fitted_width
            img.height = fitted_height
            ws.add_image(img, 'A1')
            title_start_column = 'A'
            first_title_row = 5
            ws.row_dimensions[1].height = 25
            ws.row_dimensions[2].height = 25
            ws.row_dimensions[3].height = 25
            ws.row_dimensions[4].height = 25
        except Exception:
            first_title_row = 1
            title_start_column = 'A'

    end_column_index = max(len(headers), column_index_from_string(title_start_column) + 3)
    end_column = get_column_letter(end_column_index)
    ws.merge_cells(f"{title_start_column}{first_title_row}:{end_column}{first_title_row}")
    ws[f'{title_start_column}{first_title_row}'] = title
    ws[f'{title_start_column}{first_title_row}'].font = Font(bold=True, size=14)
    ws[f'{title_start_column}{first_title_row}'].alignment = Alignment(horizontal='center')

    if not receipt_mode:
        ws.merge_cells(f"{title_start_column}{first_title_row + 1}:{end_column}{first_title_row + 1}")
        ws[f'{title_start_column}{first_title_row + 1}'] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws[f'{title_start_column}{first_title_row + 1}'].alignment = Alignment(horizontal='center')

    header_row = first_title_row + 4
    col_widths = [max(18, len(_cell_text(header)) * 1.8 + 4) for header in headers]
    for row in rows:
        for col_idx, value in enumerate(row[:len(headers)], 1):
            col_widths[col_idx - 1] = max(col_widths[col_idx - 1], max(18, len(_cell_text(value)) * 1.4 + 4))

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for row_idx, row in enumerate(rows, header_row + 1):
        for col_idx, value in enumerate(row[:len(headers)], 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = min(width, 48)

    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_report(
    title,
    headers,
    rows,
    format='pdf',
    include_signatures=False,
    signature_names=None,
    metadata_lines=None,
    receipt_mode=False,
):
    if format == 'pdf':
        return _build_pdf_response(
            title,
            headers,
            rows,
            include_signatures=include_signatures,
            signature_names=signature_names,
            metadata_lines=metadata_lines,
            receipt_mode=receipt_mode,
        )
    elif format == 'excel':
        return _build_excel_response(title, headers, rows, receipt_mode=receipt_mode)
    raise ValueError("Format must be 'pdf' or 'excel'")
